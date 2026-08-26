"""Convert the 2026 target-company workbook into JSON seed data."""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

COUNTRY_ISO = {
    "UAE": "AE",
    "United Arab Emirates": "AE",
    "Netherlands": "NL",
    "Germany": "DE",
    "Saudi Arabia": "SA",
    "Sweden": "SE",
    "Ireland": "IE",
    "Denmark": "DK",
    "Finland": "FI",
    "Norway": "NO",
}

PRIORITY_SCORE = {"S": 100, "A": 80, "B": 55, "C": 35}

# Keep previously verified ATS adapters when names match.
ATS_BY_NAME = {
    "careem": {"ats_type": "greenhouse", "board_token": "careem", "slug": "careem"},
    "n26": {"ats_type": "greenhouse", "board_token": "n26", "slug": "n26"},
    "adyen": {"ats_type": "greenhouse", "board_token": "adyen", "slug": "adyen"},
    "picnic": {"ats_type": "greenhouse", "board_token": "teampicnic", "slug": "picnic"},
    "intercom": {"ats_type": "greenhouse", "board_token": "intercom", "slug": "intercom"},
    "hubspot": {"ats_type": "greenhouse", "board_token": "hubspotjobs", "slug": "hubspot"},
    "hubspot ireland": {"ats_type": "greenhouse", "board_token": "hubspotjobs", "slug": "hubspot"},
    "personio": {"ats_type": "personio_xml", "board_token": "personio", "slug": "personio"},
    "pipedrive": {"ats_type": "lever", "board_token": "pipedrive", "slug": "pipedrive"},
    "property finder": {"slug": "property-finder"},
    "booking.com": {"slug": "booking"},
    "delivery hero": {"slug": "delivery-hero"},
}


def slugify(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return slug[:120]


def json_safe(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def likelihood_to_visa(value: str | None) -> str:
    text = (value or "").strip().lower()
    if text in {"verified/high", "high"}:
        return "likely"
    if text in {"medium/high"}:
        return "unknown"
    return "unknown"


def as_url(value) -> str | None:
    text = str(value or "").strip()
    if text.startswith(("http://", "https://")):
        return text
    return None


def ind_is_yes(value) -> bool:
    text = str(value or "").strip().lower()
    if text in {"", "n/a", "na", "no", "n", "false", "0", "none"}:
        return False
    return text in {"yes", "y", "true", "1", "recognised", "recognized", "verified sponsor"} or "sponsor" in text


def convert(xlsx_path: Path, out_path: Path) -> int:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    companies_ws = wb["Companies"]
    examples_ws = wb["Verified Examples"]
    headers = [str(cell or "").strip() for cell in next(companies_ws.iter_rows(max_row=1, values_only=True))]
    idx = {name: i for i, name in enumerate(headers)}

    examples: dict[str, list[dict]] = {}
    example_rows = list(examples_ws.iter_rows(values_only=True))
    for row in example_rows[1:]:
        if not row or not row[0]:
            continue
        examples.setdefault(str(row[0]).strip().lower(), []).append(
            {
                "claim": "international_hiring",
                "source_name": "verified_example_2026",
                "source_url": as_url(row[3] if len(row) > 3 else None),
                "type": "third_party",
                "confidence": 0.55,
                "text": row[2],
            }
        )

    records = []
    used_slugs: set[str] = set()
    for row in companies_ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx["Company"]]:
            continue
        name = str(row[idx["Company"]]).strip()
        country_name = str(row[idx["Country"]]).strip()
        if country_name not in COUNTRY_ISO:
            continue
        key = name.lower()
        overlay = ATS_BY_NAME.get(key, {})
        slug = overlay.get("slug") or slugify(name)
        base = slug
        n = 2
        while slug in used_slugs:
            slug = f"{base}-{n}"
            n += 1
        used_slugs.add(slug)

        likelihood = row[idx["Sponsorship Likelihood"]]
        ind = row[idx["IND Recognised Sponsor (NL)"]]
        evidence_text = row[idx["Visa / Relocation Evidence"]]
        last_verified = json_safe(row[idx["Last Verified"]])
        careers = as_url(row[idx["Careers URL"]])
        linkedin = as_url(row[idx["LinkedIn URL"]])

        evidence = []
        if evidence_text:
            evidence.append(
                {
                    "claim": "international_hiring",
                    "source_name": "europe_mena_target_companies_2026",
                    "source_url": careers,
                    "type": "third_party",
                    "confidence": 0.55 if likelihood_to_visa(str(likelihood)) == "likely" else 0.35,
                    "text": str(evidence_text),
                }
            )
        if ind_is_yes(ind):
            evidence.append(
                {
                    "claim": "recognized_sponsor",
                    "source_name": "IND recognised sponsor list (as recorded 2026 dataset)",
                    "source_url": careers,
                    "type": "government",
                    "confidence": 0.8,
                    "text": "Recorded as IND recognised sponsor in the 2026 research workbook.",
                }
            )
        evidence.extend(examples.get(key, []))

        records.append(
            {
                "slug": slug,
                "name": name,
                "country": COUNTRY_ISO[country_name],
                "city": str(row[idx["City"]] or "").strip() or None,
                "ats_type": overlay.get("ats_type", "career_page"),
                "board_token": overlay.get("board_token"),
                "careers_url": careers,
                "linkedin_url": linkedin,
                "priority": PRIORITY_SCORE.get(str(row[idx["Priority"]] or "A").strip(), 80),
                "enabled": True,
                "visa_status": "likely" if ind_is_yes(ind) else likelihood_to_visa(str(likelihood)),
                "extra": {
                    "industry": row[idx["Industry"]],
                    "sponsorship_likelihood": likelihood,
                    "frontend_fit": row[idx["Frontend Fit"]],
                    "fullstack_fit": row[idx["Full-stack Fit"]],
                    "stack_match": row[idx["Stack Match"]],
                    "overall_score": row[idx["Overall Score"]],
                    "priority_band": row[idx["Priority"]],
                    "last_verified": last_verified,
                    "notes": row[idx["Notes"]],
                    "ind_recognised_sponsor": ind,
                },
                "evidence": evidence,
            }
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(records, indent=2, default=str), encoding="utf-8")
    return len(records)


if __name__ == "__main__":
    import argparse

    repo = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "xlsx",
        nargs="?",
        default=r"C:\Users\USER\Downloads\europe_mena_visa_sponsorship_target_companies_2026.xlsx",
        type=Path,
    )
    parser.add_argument(
        "-o",
        "--output",
        default=repo / "backend" / "app" / "db" / "data" / "target_companies.json",
        type=Path,
    )
    args = parser.parse_args()
    count = convert(args.xlsx, args.output)
    print(f"wrote {count} companies")
